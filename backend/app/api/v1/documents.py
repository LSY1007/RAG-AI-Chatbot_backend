from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status, Query
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.core.security import decode_access_token
from app.models.user import User
from app.models.chat import Document, ChatRoom
from app.services.rag_service import process_document
from app.services.vector_service import delete_documents_by_room
import pdfplumber
import docx
import re
import io

router = APIRouter()
security = HTTPBearer()

def get_current_user(
    credentials: HTTPAuthorizationCredentials = Depends(security),
    db: Session = Depends(get_db)
) -> User:
    payload = decode_access_token(credentials.credentials)
    if not payload:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="유효하지 않은 토큰입니다.")
    user = db.query(User).filter(User.id == int(payload.get("sub"))).first()
    if not user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="유저를 찾을 수 없습니다.")
    return user

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    texts = []
    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    texts.append(cell.text.strip())
    full_xml = doc.element.xml
    found = re.findall(r'<w:t[^>]*>([^<]+)</w:t>', full_xml)
    for t in found:
        t = t.strip()
        if t and t not in texts:
            texts.append(t)
    return "\n".join(texts)

def extract_text_from_xlsx(file_bytes: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        texts = []
        for sheet in wb.worksheets:
            texts.append(f"[시트: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_texts = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                if row_texts:
                    texts.append("\t".join(row_texts))
        return "\n".join(texts)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Excel 파일 처리 오류: {str(e)}")

def extract_text_from_pptx(file_bytes: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_bytes))
        texts = []
        for i, slide in enumerate(prs.slides):
            texts.append(f"[슬라이드 {i+1}]")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
        return "\n".join(texts)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"PPT 파일 처리 오류: {str(e)}")

def extract_text_from_hwp(file_bytes: bytes) -> str:
    try:
        import olefile
        if not olefile.isOleFile(io.BytesIO(file_bytes)):
            raise HTTPException(status_code=400, detail="유효하지 않은 HWP 파일입니다.")
        ole = olefile.OleFileIO(io.BytesIO(file_bytes))
        texts = []
        if ole.exists('PrvText'):
            text_data = ole.openstream('PrvText').read()
            texts.append(text_data.decode('utf-16-le', errors='ignore'))
        ole.close()
        return "\n".join(texts)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"HWP 파일 처리 오류: {str(e)}")

def extract_text(file_bytes: bytes, filename: str) -> str:
    filename_lower = filename.lower()
    if filename_lower.endswith(".pdf"):
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            return "\n".join([page.extract_text() or "" for page in pdf.pages])
    elif filename_lower.endswith(".txt"):
        return file_bytes.decode("utf-8")
    elif filename_lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif filename_lower.endswith((".xlsx", ".xls")):
        return extract_text_from_xlsx(file_bytes)
    elif filename_lower.endswith((".pptx", ".ppt")):
        return extract_text_from_pptx(file_bytes)
    elif filename_lower.endswith(".hwp"):
        return extract_text_from_hwp(file_bytes)
    else:
        raise HTTPException(status_code=400, detail="PDF, TXT, DOCX, XLSX, PPTX, HWP 파일만 지원합니다.")


@router.post("/upload", summary="문서 업로드")
async def upload_document(
    file: UploadFile = File(...),
    room_id: int = Query(..., description="채팅방 ID"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # 채팅방 존재 확인
    room = db.query(ChatRoom).filter(
        ChatRoom.id == room_id,
        ChatRoom.user_id == current_user.id
    ).first()
    if not room:
        raise HTTPException(status_code=404, detail="채팅방을 찾을 수 없습니다.")

    file_bytes = await file.read()
    filename = file.filename

    text = extract_text(file_bytes, filename)
    if not text.strip():
        raise HTTPException(status_code=400, detail="텍스트를 추출할 수 없습니다.")

    chunk_count = process_document(
        text=text,
        filename=filename,
        user_id=current_user.id,
        room_id=room_id
    )

    doc = Document(
        user_id=current_user.id,
        room_id=room_id,
        filename=filename,
        file_type=filename.split(".")[-1].lower(),
        chunk_count=chunk_count
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    return {
        "message": "문서 업로드 완료",
        "filename": filename,
        "chunk_count": chunk_count,
        "document_id": doc.id
    }


@router.delete("/{document_id}", summary="문서 삭제")
def delete_document(
    document_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(Document).filter(
        Document.id == document_id,
        Document.user_id == current_user.id
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="문서를 찾을 수 없습니다.")

    # ChromaDB에서 삭제
    if doc.room_id:
        delete_documents_by_room(
            user_id=current_user.id,
            room_id=doc.room_id,
            filename=doc.filename
        )

    db.delete(doc)
    db.commit()
    return {"message": "문서가 삭제되었습니다."}


@router.get("/", summary="문서 목록")
def get_documents(
    room_id: int = Query(..., description="채팅방 ID"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """채팅방별 문서 목록 조회"""
    docs = db.query(Document).filter(
        Document.user_id == current_user.id,
        Document.room_id == room_id
    ).order_by(Document.created_at.desc()).all()

    return [
        {
            "id": d.id,
            "filename": d.filename,
            "file_type": d.file_type,
            "chunk_count": d.chunk_count,
            "created_at": str(d.created_at)
        }
        for d in docs
    ]
